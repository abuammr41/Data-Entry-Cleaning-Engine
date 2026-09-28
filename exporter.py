import math
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
BAD_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
ISSUE_ROW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)


def _safe(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)[:32000]
    return v


def _sheet(wb, title, df, bad=None):
    ws = wb.create_sheet(title)
    ws.append(list(df.columns))
    for c in ws[1]:
        c.fill, c.font, c.border = HEADER_FILL, Font(bold=True, color="FFFFFF"), BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    cols = list(df.columns)
    for r_i, (idx, row) in enumerate(df.iterrows(), start=2):
        ws.append([_safe(v) for v in row.tolist()])
        has_issue = row.get("Issues", "OK") != "OK"
        for c_i, col in enumerate(cols, start=1):
            cell = ws.cell(r_i, c_i)
            cell.border = BORDER
            cell.font = Font(size=10)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
            cell.alignment = Alignment(horizontal="right" if isinstance(cell.value, (int, float)) else "left")
            if bad is not None and col in bad.columns and bad.at[idx, col]:
                cell.fill = BAD_FILL
            elif has_issue and col == "Issues":
                cell.fill = ISSUE_ROW_FILL
    for i, col in enumerate(cols, 1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).head(500)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 3, 12), 50)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def export_formatted_excel(df, bad, report, filename="Cleaned_Data_Entry_Output.xlsx"):
    output_file = filename if os.path.isabs(filename) else os.path.join(OUTPUT_DIR, filename)
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Summary")
    rows = [
        ("Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Rows received", report["rows_in"]),
        ("Empty rows removed", report["empty_rows_removed"]),
        ("Duplicate rows removed", report["duplicates_removed"]),
        ("Clean rows delivered", report["rows_out"]),
        ("Rows needing review", report["rows_with_issues"]),
        ("", ""),
        ("FIELD CHECKS", "Invalid / Missing"),
    ]
    for kind, c in report["counts"].items():
        if any(t == kind for t in report["column_types"].values()):
            rows.append((kind.upper(), f"{c['invalid']} / {c['missing']}"))
    for label, value in rows:
        ws.append([label, value])
    for c in ws["A"]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22

    _sheet(wb, "Clean_Data", df, bad)
    review = df[df["Issues"] != "OK"]
    if len(review):
        _sheet(wb, "Needs_Review", review, bad.loc[review.index])

    try:
        wb.save(output_file)
    except PermissionError:
        base, ext = os.path.splitext(output_file)
        output_file = f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
        wb.save(output_file)
        print("NOTE: The previous file was open in Excel, saved with a new name.")
    return output_file
